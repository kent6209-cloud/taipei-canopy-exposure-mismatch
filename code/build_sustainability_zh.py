# -*- coding: utf-8 -*-
"""Convert 論文_全文_v15.docx into MDPI Sustainability (Chinese edition) layout.

Input   Paper/論文_全文_v15.docx
Template E:\\SCI\\new\\SCI_Ready\\sustainability-4529713.docx  (39 MDPI styles, A4,
          margins 12.7/12.7/25/16 mm, three-line table style)
Output  Paper/論文_全文_v15_Sustainability中文版.docx

Mapping
  Heading 1 (中文題目)            -> MDPI_1.2_title
  英文題目                        -> MDPI_1.2_title (12 pt, 非粗體)
  摘要 / Abstract (+ 內文)        -> MDPI_1.7_abstract
  關鍵詞 / Keywords               -> MDPI_1.8_keywords
  目錄                            -> 略去（期刊稿無目錄）
  Heading 3 (1.1, 2.3, 附錄 A…)   -> MDPI_2.2_heading2
  Heading 2 (參考文獻 / 附錄)      -> MDPI_2.1_heading1
  內文                            -> MDPI_3.1_text（清單 -> MDPI_3.7_itemize）
  表標題 / 表格                    -> MDPI_4.1_table_caption / MDPI_4.1_three_line_table
  圖 / 圖說                        -> MDPI_5.2_figure / MDPI_5.1_figure_caption
  參考文獻條目                     -> MDPI_8.1_references
  中文字型：以 eastAsia='PMingLiU'（新細明體）補齊，拉丁字保留模板字型（Palatino Linotype）
"""
import re
import shutil
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph

ROOT = Path(r'E:\GeoAI\20260909_taiwan_chmv2\Paper')
SRC = ROOT / '論文_全文_v15.docx'
TPL = Path(r'E:\SCI\new\SCI_Ready\sustainability-4529713.docx')
CANON_SUFFIX = 'v15'
OUTDIR = None


def out_paths(suffix=None):
    """Output paths for the journal build; suffix follows CANON_SUFFIX / --v16."""
    s = suffix or CANON_SUFFIX
    return {
        'full': ROOT / ('論文_全文_{}_Sustainability中文版.docx'.format(s)),
        'main': ROOT / ('論文_全文_{}_Sustainability中文版_主文.docx'.format(s)),
        'supp': ROOT / ('論文_全文_{}_Sustainability中文版_附錄（補充材料）.docx'.format(s)),
    }


_paths = out_paths()
OUT = _paths['full']
OUT_MAIN = _paths['main']
OUT_SUPP = _paths['supp']
CHAPTER_TITLES = {1: '緒論', 2: '文獻回顧', 3: '理論框架', 4: '研究設計與方法',
                  5: '研究結果', 6: '討論', 7: '結論與摘要'}
EA_FONT = 'PMingLiU'
LATIN_FONT = 'Palatino Linotype'

T_CAPTION = re.compile(r'^\**表\s')
F_CAPTION = re.compile(
    r'^(\**圖\s*[A-Z0-9]+(?:-\d+)?\u3000'          # 圖 5-1　標題
    r'|\**附錄圖\s*[A-Z0-9]+-\d+\u3000'            # 附錄圖 A-1　標題（映射前）
    r'|\*?Figure\s+S?\d+(?:-\d+)?[.\u3000]'       # Figure 5-1. / Figure S1　（映射後）
    r'|\*?Appendix Figure\s+[A-Z]?-?\d+\.)')        # Appendix Figure A-1.

REF_LIKE = re.compile(r'^\**\d{1,3}\.\s')


# --- Supplementary numbering (P0-7) -----------------------------------------
SUPP_MAP = [
    ('附錄圖 A-1', 'Figure S1'), ('附錄圖 A-2', 'Figure S2'), ('附錄圖 A-6', 'Figure S3'),
    ('附錄圖 A-3', 'Figure S4'), ('附錄圖 A-4', 'Figure S5'), ('附錄圖 A-5', 'Figure S6'),
    ('附錄 G-2', 'Supplementary S8'), ('附錄 A', 'Supplementary S1'), ('附錄 B', 'Supplementary S2'),
    ('附錄 C', 'Supplementary S3'), ('附錄 D', 'Supplementary S4'), ('附錄 E', 'Supplementary S5'),
    ('附錄 F', 'Supplementary S6'), ('附錄 G', 'Supplementary S7'), ('附錄 H', 'Supplementary S9'),
    ('附錄 I', 'Supplementary S10'), ('附錄 J', 'Supplementary S11'),
]
SUPP_INDEX = [
    ('Supplementary S1', '資料來源、空間基線與可重現性（Data sources, spatial baseline and reproducibility）'),
    ('Supplementary S2', '可重現程式與執行稽核（Reproducible code base and software audit）'),
    ('Supplementary S3', '空間自相關診斷與空間有效推論（Moran\'s I, LISA, block inference）'),
    ('Supplementary S4', '案例編碼協定（Coding protocols for community and corporate projects）'),
    ('Supplementary S5', '不確定性傳播與 CHMv2 樹冠高度驗證（Uncertainty propagation; Figure S1）'),
    ('Supplementary S6', '圖版索引（Figure index and submission mapping）'),
    ('Supplementary S7', '全臺鄉鎮市區外部對照（National township comparison; Figure S2）'),
    ('Supplementary S8', '全臺 22 縣市對照與交叉檢核（County comparison; Figure S3）'),
    ('Supplementary S9', '官方山坡地與使用分區交叉檢核（Slope land and zoning; Figure S4）'),
    ('Supplementary S10', '年 NPP 通量產品可用性檢核（Flux-product usability audit; Figure S5）'),
    ('Supplementary S11', '行道樹結構分解與區級樹冠結構（Street-tree structure; Figure S6）'),
]


# --- Journal numbering (Figures / Tables) -----------------------------------
FIG_MAP = [
    ('附錄圖 A-1', 'Figure S1'), ('附錄圖 A-2', 'Figure S2'), ('附錄圖 A-6', 'Figure S3'),
    ('附錄圖 A-3', 'Figure S4'), ('附錄圖 A-4', 'Figure S5'), ('附錄圖 A-5', 'Figure S6'),
    ('Appendix Figure A-1.', 'Figure S1.'), ('Appendix Figure A-2.', 'Figure S2.'),
    ('Appendix Figure A-3.', 'Figure S4.'), ('Appendix Figure A-4.', 'Figure S5.'),
    ('Appendix Figure A-5.', 'Figure S6.'), ('Appendix Figure A-6.', 'Figure S3.'),
    ('Figure A-1', 'Figure S1'), ('Figure A-2', 'Figure S2'),
    ('圖 2-1', 'Figure 1'), ('圖 3-1', 'Figure 2'), ('圖 4-1', 'Figure 3'), ('圖 4-2', 'Figure 4'),
    ('圖 4-3', 'Figure 5'), ('圖 5-1', 'Figure 6'), ('圖 5-2', 'Figure 7'), ('圖 5-3', 'Figure 8'),
    ('圖 5-4', 'Figure 9'), ('圖 5-5', 'Figure 10'),
    ('Figure 2-1', 'Figure 1'), ('Figure 3-1', 'Figure 2'), ('Figure 4-1', 'Figure 3'),
    ('Figure 4-2', 'Figure 4'), ('Figure 4-3', 'Figure 5'), ('Figure 5-1', 'Figure 6'),
    ('Figure 5-2', 'Figure 7'), ('Figure 5-3', 'Figure 8'), ('Figure 5-4', 'Figure 9'),
    ('Figure 5-5', 'Figure 10'),
]
TABLE_MAP = [
    ('表 S10', 'Table S10'), ('表 S9', 'Table S9'), ('表 S8', 'Table S8'), ('表 S7', 'Table S7'),
    ('表 S6', 'Table S6'), ('表 S5', 'Table S5'), ('表 S4', 'Table S4'), ('表 S3', 'Table S3'),
    ('表 S2', 'Table S2'), ('表 S1', 'Table S1'),
    ('表 1.1', 'Table 1'), ('表 1.2', 'Table 2'), ('表 3.1', 'Table 3'), ('表 3.2', 'Table 4'),
    ('表 3.3', 'Table 5'), ('表 4.1', 'Table 6'), ('表 4.2', 'Table 7'), ('表 4.3', 'Table 8'),
    ('表 4.4', 'Table 9'), ('表 4.5', 'Table 10'), ('表 4.6', 'Table 11'), ('表 4.7', 'Table 12'),
    ('表 4.8', 'Table 13'), ('表 5.1', 'Table 14'), ('表 5.2', 'Table 15'), ('表 5.3', 'Table 16'),
    ('表 5.4', 'Table 17'), ('表 5.5', 'Table 18'), ('表 5.6', 'Table 19'), ('表 5.7', 'Table 20'),
    ('表 5.8', 'Table 21'), ('表 5.9', 'Table 22'),
]
BACK_MATTER = [
    ('Author Contributions: Conceptualization, methodology, formal analysis, data curation, '
     'writing—original draft preparation and writing—review and editing were performed by the '
     'author(s). All author(s) have read and agreed to the published version of the manuscript.'),
    ('Funding: This research received no external funding.'),
    ('Data Availability Statement: All datasets used in this study are publicly available; the '
     'sources and processing rules are listed in Table S1 and Table S2. Derived spatial layers '
     'and analysis scripts are archived at Zenodo: https://doi.org/10.5281/zenodo.22865864 '
     '(version v1.0.0; concept DOI 10.5281/zenodo.22865863; MIT licence for code, CC BY 4.0 for '
     'derived results).'),
    ('Acknowledgments: The author(s) thank the Taipei City Government (Parks and Street Lights '
     'Office; Urban Development Bureau) and the Forestry and Nature Conservation Agency for '
     'publicly releasing the datasets used in this study.'),
    ('Conflicts of Interest: The author(s) declare no conflicts of interest.'),
]
REF_SUBHEADS = ('學術文獻', '法規與政策文件', '資料來源')

CURRENT_MODE = 'full'


def map_refs(text, mode=None):
    """Journal numbering: appendix -> Supplementary S#, chapters -> Figure N / Table N.

    Uses the module-level CURRENT_MODE (set per build) so that call sites passing the
    build argument cannot override the outdir-driven 'main' mapping.
    """
    if CURRENT_MODE == 'full':
        return text
    for a, b in SUPP_MAP + FIG_MAP + TABLE_MAP:
        text = text.replace(a, b)
    return text


def style_run(run, size=None, bold=None, italic=None):
    run.font.name = LATIN_FONT
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = rpr.makeelement(qn('w:rFonts'), {})
        rpr.insert(0, rf)
    rf.set(qn('w:ascii'), LATIN_FONT)
    rf.set(qn('w:hAnsi'), LATIN_FONT)
    rf.set(qn('w:eastAsia'), EA_FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_para(doc, text, style, size=None, bold=None, italic=None, align=None, mode=None):
    p = doc.add_paragraph(style=style)
    r = p.add_run(map_refs(text, mode or CURRENT_MODE))
    style_run(r, size=size, bold=bold, italic=italic)
    if align:
        p.alignment = align
    return p


def copy_images(src_para, doc, out_dir):
    """把來源段落中的圖片複製到新文件，回傳張數。"""
    blips = src_para._p.findall('.//' + qn('a:blip'))
    n = 0
    for blip in blips:
        rid = blip.get(qn('r:embed'))
        if not rid:
            continue
        part = src_para.part.related_parts[rid]
        tmp = out_dir / ('_img%d%s' % (n, Path(str(part.partname)).suffix or '.png'))
        tmp.write_bytes(part.blob)
        p = doc.add_paragraph(style='MDPI_5.2_figure')
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(tmp), width=Cm(15.0))
        n += 1
    return n


def new_table(doc, src_tbl, mode=None):
    t = doc.add_table(rows=len(src_tbl.rows), cols=len(src_tbl.columns))
    t.style = doc.styles['MDPI_4.1_three_line_table']
    for i, row in enumerate(src_tbl.rows):
        for j, cell in enumerate(row.cells):
            txt = map_refs(cell.text.strip(), mode or CURRENT_MODE)
            dst = t.cell(i, j)
            dst.text = ''
            r = dst.paragraphs[0].add_run(txt)
            style_run(r, size=9, bold=(i == 0))
    return t


def build(mode='full', outdir=None):
    """mode: full | main | supp"""
    global CURRENT_MODE, OUTDIR
    OUTDIR = outdir
    # 若輸出至自訂資料夾（投稿交付），全文件亦改用投稿編號（Figure/Table N、S#）
    CURRENT_MODE = 'main' if (mode == 'full' and outdir is not None) else mode
    if not SRC.exists():
        print('missing', SRC, file=sys.stderr)
        return 1
    src = Document(str(SRC))
    doc = Document(str(TPL))
    # 清空模板內容，保留樣式與版面設定
    body = doc.element.body
    for child in list(body.iterchildren()):
        if child.tag in (qn('w:p'), qn('w:tbl')):
            body.remove(child)

    tmp_dir = Path(tempfile.mkdtemp(prefix='mdpizh_'))
    state = {'zone': 'front', 'skip_toc': False, 'chapter': 0, 'captioned': False,
             'appendix_started': False, 'ref_no': 0}
    n_tbl = n_img = n_cap = 0

    children = list(src.element.body.iterchildren())
    paras_done = 0
    for child in children:
        if child.tag == qn('w:tbl'):
            if mode == 'supp' and not state['appendix_started']:
                state['captioned'] = False
                continue
            if not state['captioned']:
                add_para(doc, '表（表號與表題待補）【待補】', 'MDPI_4.1_table_caption', mode=mode)
            new_table(doc, Table(child, src), mode)
            state['captioned'] = False
            n_tbl += 1
            continue
        if child.tag != qn('w:p'):
            continue
        p = Paragraph(child, src)
        text = p.text.strip()
        sn = p.style.name
        paras_done += 1

        # --- 目錄：略去 ---
        if sn == 'Heading 2' and text == '目錄':
            state['skip_toc'] = True
            continue
        if state['skip_toc']:
            if sn.startswith('Heading'):
                state['skip_toc'] = False
                state['zone'] = 'body'
            else:
                continue

        # --- 主文／補充材料之切分：附錄起為補充材料 ---
        if sn == 'Heading 2' and text == '附錄':
            state['appendix_started'] = True
            if mode == 'main':
                break
            if mode == 'supp':
                add_para(doc, 'Supplementary Materials（補充材料）', 'MDPI_2.1_heading1')
                add_para(doc, '本補充材料以區段 S1–S11 編號；圖版另以 Figure S1–S6 連續編號，表格以 Table S1–S10 連續編號，主文之引註已同步對應。',
                         'MDPI_3.1_text', mode=mode)
                add_para(doc, '| 編號 | 內容 |', 'MDPI_4.1_table_caption', mode=mode)
                for sid, desc in SUPP_INDEX:
                    add_para(doc, '| {} | {} |'.format(sid, desc), 'MDPI_3.1_text', mode=mode)
        if mode == 'supp' and not state['appendix_started']:
            continue

        # --- 圖片 ---
        if p._p.findall('.//' + qn('a:blip')):
            n_img += copy_images(p, doc, tmp_dir)
            continue

        if not text:
            continue

        # --- 標題層級 ---
        if sn == 'Heading 1':
            add_para(doc, text, 'MDPI_1.2_title', mode=mode)
            continue
        if sn == 'Heading 2':
            if text in ('摘要', 'Abstract'):
                add_para(doc, text, 'MDPI_1.7_abstract', bold=True)
                state['zone'] = 'abstract'
                continue
            if text == '參考文獻':
                if mode != 'supp':
                    for para in BACK_MATTER:
                        add_para(doc, para, 'MDPI_3.1_text', mode=mode)
                state['zone'] = 'refs'
                add_para(doc, text, 'MDPI_2.1_heading1', mode=mode)
                continue
            if text == '附錄':
                state['zone'] = 'appendix'
                if mode != 'supp':
                    add_para(doc, text, 'MDPI_2.1_heading1', mode=mode)
                continue
            state['zone'] = 'body'
            add_para(doc, text, 'MDPI_2.1_heading1', mode=mode)
            continue
        if sn == 'Heading 3':
            if state['zone'] not in ('refs', 'appendix'):
                state['zone'] = 'body'
            m = re.match(r'^(\d+)\.', text)
            if m and state['zone'] == 'body':
                ch = int(m.group(1))
                if ch != state['chapter'] and ch in CHAPTER_TITLES:
                    state['chapter'] = ch
                    add_para(doc, '{}. {}'.format(ch, CHAPTER_TITLES[ch]), 'MDPI_2.1_heading1')
            add_para(doc, text, 'MDPI_2.2_heading2', mode=mode)
            continue

        # --- 前置：英文題目、作者與定位句 ---
        if paras_done == 2:
            add_para(doc, text, 'MDPI_1.2_title', size=12, bold=False,
                     align=WD_ALIGN_PARAGRAPH.CENTER)
            add_para(doc, '作者姓名【待補】', 'MDPI_1.3_authornames', align=WD_ALIGN_PARAGRAPH.CENTER, mode=mode)
            add_para(doc, '服務單位【待補】（通訊作者：E-mail【待補】）', 'MDPI_1.6_affiliation',
                     align=WD_ALIGN_PARAGRAPH.CENTER)
            add_para(doc, '文章類型：研究論文', 'MDPI_1.1_article_type', mode=mode)
            continue
        if paras_done == 3:
            add_para(doc, text, 'MDPI_1.6_affiliation', italic=True, mode=mode)
            continue

        # --- 摘要區 ---
        if state['zone'] == 'abstract':
            if text.startswith('關鍵詞') or text.startswith('Keywords'):
                add_para(doc, text, 'MDPI_1.8_keywords', mode=mode)
            else:
                add_para(doc, text, 'MDPI_1.7_abstract', mode=mode)
            continue

        # --- 參考文獻 ---
        if state['zone'] == 'refs':
            if text in REF_SUBHEADS:
                continue
            state['ref_no'] += 1
            body = re.sub(r'^\s*\d+\.\s*', '', text)
            add_para(doc, '[{}] {}'.format(state['ref_no'], body), 'MDPI_8.1_references', mode=mode)
            continue

        # --- 表標題 / 圖說 ---
        if T_CAPTION.match(text):
            add_para(doc, text, 'MDPI_4.1_table_caption', mode=mode)
            state['captioned'] = True
            n_cap += 1
            continue
        if F_CAPTION.match(text):
            add_para(doc, text, 'MDPI_5.1_figure_caption', mode=mode)
            n_cap += 1
            continue

        # --- 內文 / 清單 / 引文 ---
        if sn == 'Compact':
            add_para(doc, '•\t' + text, 'MDPI_3.7_itemize', mode=mode)
        elif sn == 'Block Text':
            p2 = add_para(doc, text, 'MDPI_3.3_text_space_after', italic=True, mode=mode)
            p2.paragraph_format.left_indent = Cm(0.6)
        else:
            add_para(doc, text, 'MDPI_3.1_text', mode=mode)

    target = {'full': OUT, 'main': OUT_MAIN, 'supp': OUT_SUPP}[mode]
    if OUTDIR:
        OUTDIR.mkdir(parents=True, exist_ok=True)
        target = OUTDIR / target.name
    doc.save(str(target))
    shutil.rmtree(tmp_dir, ignore_errors=True)
    print('wrote {} ({:,} bytes)  tables={} figures={} captions={}'.format(
        target.name, target.stat().st_size, n_tbl, n_img, n_cap))
    return 0


def main():
    if '--v16' in sys.argv:
        global CANON_SUFFIX, OUT, OUT_MAIN, OUT_SUPP
        CANON_SUFFIX = 'v16'
        _p = out_paths('v16')
        OUT, OUT_MAIN, OUT_SUPP = _p['full'], _p['main'], _p['supp']
    outdir = None
    for i, a in enumerate(sys.argv):
        if a == '--outdir' and i + 1 < len(sys.argv):
            outdir = Path(sys.argv[i + 1])
    modes = ['full', 'main', 'supp'] if '--split' in sys.argv else ['full']
    rc = 0
    for m in modes:
        rc = build(m, outdir) or rc
    return rc


if __name__ == '__main__':
    sys.exit(main())
